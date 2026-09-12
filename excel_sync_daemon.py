import time
import requests
import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os
import csv

DB_PATH = r"C:\Users\rohan\OneDrive\Desktop\YourInteriorDesk\Client_Designer_DB\client_designer.db"
EXCEL_PATH = r"C:\Users\rohan\OneDrive\Desktop\YourInteriorDesk\Client_Designer_DB\Client_Designer.xlsx"
CSV_PATH = r"C:\Users\rohan\OneDrive\Desktop\YourInteriorDesk\Client_Designer_DB\Client_Designer.csv"

# Cloud Vercel Endpoints or Local Server Endpoints
VERCEL_API = os.environ.get("VERCEL_API_URL", "https://your-interior-desk.vercel.app/api/submissions")
VERCEL_REVIEWS_API = os.environ.get("VERCEL_REVIEWS_API_URL", "https://your-interior-desk.vercel.app/api/reviews")
GITHUB_SUBS_API = "https://raw.githubusercontent.com/mono-xyxy/your-interior-desk/main/data/submissions.json"
LOCAL_API = "http://localhost:3000/api/submissions"
LOCAL_REVIEWS_API = "http://localhost:3000/api/reviews"

# Excel Beautification Styles (Steel Navy Header & Clean Typography)
header_fill = PatternFill(start_color='101B2E', end_color='101B2E', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='F8FAFC')

row_fill_even = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
row_fill_odd = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_wrap = Alignment(horizontal='left', vertical='top', wrap_text=True)

def auto_fit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 6, 20), 70)

def apply_beautification(ws):
    ws.row_dimensions[1].height = 26
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = row_fill_even if row_idx % 2 == 0 else row_fill_odd
        ws.row_dimensions[row_idx].height = 24
        for cell in row:
            cell.fill = fill
            cell.border = thin_border
            cell.font = Font(name='Calibri', size=11, color='0F172A')
            
            if cell.column in [1, 2, 3, 5, 6, 7, 8, 10, 11]:
                cell.alignment = align_center
            elif cell.column in [12]:
                cell.alignment = align_wrap
            else:
                cell.alignment = align_left

    auto_fit_columns(ws)

def fetch_cloud_submissions():
    merged = {}
    for endpoint in [VERCEL_API, GITHUB_SUBS_API, LOCAL_API]:
        try:
            res = requests.get(endpoint, timeout=3)
            if res.status_code == 200:
                data = res.json()
                if isinstance(data, list):
                    items = data
                else:
                    items = data.get("submissions", [])
                for item in items:
                    if isinstance(item, dict) and item.get("id"):
                        merged[item["id"]] = item
        except Exception:
            pass
    return list(merged.values())

def fetch_cloud_reviews():
    merged = {}
    for endpoint in [VERCEL_REVIEWS_API, LOCAL_REVIEWS_API]:
        try:
            res = requests.get(endpoint, timeout=3)
            if res.status_code == 200:
                data = res.json()
                if isinstance(data, list):
                    items = data
                else:
                    items = data.get("reviews", [])
                for item in items:
                    if isinstance(item, dict) and item.get("id"):
                        merged[item["id"]] = item
        except Exception:
            pass
    return list(merged.values())

def get_all_db_submissions():
    if os.path.exists(DB_PATH):
        try:
            conn = sqlite3.connect(DB_PATH, timeout=10.0)
            c = conn.cursor()
            c.execute("""
                SELECT id, timestamp, role, fullName, email, phone, location, budget, 
                       socialHandles, description, wordCount, signed_status, signatureFullName 
                FROM submissions ORDER BY timestamp ASC
            """)
            rows = c.fetchall()
            conn.close()
            return [
                {
                    "id": r[0], "timestamp": r[1], "role": r[2], "fullName": r[3],
                    "email": r[4], "phone": r[5], "location": r[6], "budget": r[7],
                    "socialHandles": r[8], "description": r[9], "wordCount": r[10],
                    "signed_status": r[11] or "signed", "signatureFullName": r[12] or ""
                } for r in rows
            ]
        except Exception:
            pass
    return []

def get_all_db_reviews():
    if os.path.exists(DB_PATH):
        try:
            conn = sqlite3.connect(DB_PATH, timeout=10.0)
            c = conn.cursor()
            c.execute("SELECT id, timestamp, name, email, role, ratingScore, ratingKeyword, emoji, reviewText, matchedKeywords, wordCount FROM reviews ORDER BY timestamp ASC")
            rows = c.fetchall()
            conn.close()
            return [
                {
                    "id": r[0], "timestamp": r[1], "name": r[2], "email": r[3],
                    "role": r[4], "ratingScore": r[5], "ratingKeyword": r[6],
                    "emoji": r[7], "reviewText": r[8], "matchedKeywords": r[9], "wordCount": r[10]
                } for r in rows
            ]
        except Exception:
            pass
    return []

def save_to_sqlite(subs, reviews):
    if not os.path.exists(os.path.dirname(DB_PATH)):
        os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    try:
        conn = sqlite3.connect(DB_PATH, timeout=10.0)
        c = conn.cursor()
        c.execute("PRAGMA busy_timeout = 10000")
        c.execute("""
            CREATE TABLE IF NOT EXISTS submissions (
                id TEXT PRIMARY KEY,
                timestamp TEXT,
                role TEXT,
                fullName TEXT,
                email TEXT,
                phone TEXT,
                location TEXT,
                budget TEXT,
                socialHandles TEXT,
                description TEXT,
                wordCount INTEGER,
                signed_status TEXT DEFAULT 'signed',
                signatureFullName TEXT,
                signatureFileName TEXT,
                signatureFilePath TEXT,
                termsAccepted INTEGER DEFAULT 1,
                synced_to_excel INTEGER DEFAULT 0
            )
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS reviews (
                id TEXT PRIMARY KEY,
                timestamp TEXT,
                name TEXT,
                email TEXT,
                role TEXT,
                ratingScore INTEGER,
                ratingKeyword TEXT,
                emoji TEXT,
                reviewText TEXT,
                matchedKeywords TEXT,
                wordCount INTEGER
            )
        """)
        for col_def in [
            "ALTER TABLE submissions ADD COLUMN socialHandles TEXT",
            "ALTER TABLE submissions ADD COLUMN signed_status TEXT DEFAULT 'signed'",
            "ALTER TABLE submissions ADD COLUMN signatureFullName TEXT",
            "ALTER TABLE submissions ADD COLUMN signatureFileName TEXT",
            "ALTER TABLE submissions ADD COLUMN signatureFilePath TEXT",
            "ALTER TABLE submissions ADD COLUMN termsAccepted INTEGER DEFAULT 1"
        ]:
            try:
                c.execute(col_def)
            except Exception:
                pass

        for s in subs:
            role_str = "Designer" if s.get("role", "").lower() == "designer" else "Client"
            c.execute("""
                INSERT OR REPLACE INTO submissions 
                (id, timestamp, role, fullName, email, phone, location, budget, socialHandles, description, wordCount, signed_status, signatureFullName, synced_to_excel)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1)
            """, (
                s.get("id"), s.get("timestamp"), role_str, s.get("fullName"),
                s.get("email"), s.get("phone"), s.get("location"), s.get("budget"),
                s.get("socialHandles", ""), s.get("description"), s.get("wordCount", 0),
                s.get("signedStatus") or s.get("signed_status") or "signed",
                s.get("signatureFullName", "")
            ))

        for r in reviews:
            c.execute("""
                INSERT OR REPLACE INTO reviews 
                (id, timestamp, name, email, role, ratingScore, ratingKeyword, emoji, reviewText, matchedKeywords, wordCount)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                r.get("id"), r.get("timestamp"), r.get("name"), r.get("email"),
                r.get("role"), r.get("ratingScore", 5), r.get("ratingKeyword", "Good"),
                r.get("emoji", "😊"), r.get("reviewText"), r.get("matchedKeywords", ""), r.get("wordCount", 0)
            ))

        conn.commit()
        conn.close()
    except Exception:
        pass

def sync_1second():
    # 1. Fetch Cloud API & GitHub submissions & reviews
    cloud_subs = fetch_cloud_submissions()
    cloud_revs = fetch_cloud_reviews()

    # 2. Save Cloud items into local SQLite DB
    if cloud_subs or cloud_revs:
        save_to_sqlite(cloud_subs, cloud_revs)

    # 3. Read ALL cumulative records from SQLite DB
    subs = get_all_db_submissions()
    reviews = get_all_db_reviews()

    if not subs and not reviews:
        return

    dir_path = os.path.dirname(EXCEL_PATH)
    os.makedirs(dir_path, exist_ok=True)

    # 4. Update CSV Log Files with ALL cumulative records
    try:
        with open(CSV_PATH, mode='w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow([
                'Submission Id', 'Timestamp', 'Role', 'Full Name', 'Email Address', 
                'Phone Number', 'Location', 'Budget Range (₹)', 'Social Handles', 
                'Signed Status', 'Signature Name', 'Word Count', 'Project Description'
            ])
            for s in reversed(subs):
                role_title = "Designer" if s.get("role", "").lower() == "designer" else "Client"
                writer.writerow([
                    s.get("id"), s.get("timestamp"), role_title,
                    s.get("fullName"), s.get("email"), s.get("phone"),
                    s.get("location"), s.get("budget"), s.get("socialHandles", ""),
                    s.get("signed_status", "signed"), s.get("signatureFullName", ""),
                    s.get("wordCount", 0), s.get("description")
                ])
    except Exception:
        pass

    try:
        reviews_csv = os.path.join(dir_path, "Client_Designer_Reviews.csv")
        with open(reviews_csv, mode='w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(['Review Id', 'Timestamp', 'Author Name', 'Email Address', 'Role', 'Sentiment Keyword', 'Emoji', 'Rating Score', 'Matched Keywords', 'Word Count', 'Review Text'])
            for r in reversed(reviews):
                writer.writerow([
                    r.get("id"), r.get("timestamp"), r.get("name"), r.get("email"),
                    r.get("role"), r.get("ratingKeyword"), r.get("emoji"),
                    r.get("ratingScore"), r.get("matchedKeywords"), r.get("wordCount"), r.get("reviewText")
                ])
    except Exception:
        pass

    # 5. Update Excel Workbook (.xlsx) with ALL cumulative records
    try:
        wb = openpyxl.Workbook()

        # Sheet 1: Designers
        ws_designers = wb.active
        ws_designers.title = 'Designers'
        ws_designers.append([
            'Submission Id', 'Timestamp', 'Full Name', 'Email Address', 'Phone Number', 
            'Working Location', 'Budget Fee (₹)', 'Signed Status', 'Signature Name', 
            'Word Count', 'Social Handles', 'Professional Overview'
        ])

        # Sheet 2: Clients
        ws_clients = wb.create_sheet(title='Clients')
        ws_clients.append([
            'Submission Id', 'Timestamp', 'Full Name', 'Email Address', 'Phone Number', 
            'Property Location', 'Offered Budget (₹)', 'Signed Status', 'Signature Name', 
            'Word Count', 'Project Scope'
        ])

        # Sheet 3: Master Log
        ws_master = wb.create_sheet(title='All Submissions')
        ws_master.append([
            'Submission Id', 'Timestamp', 'Role Type', 'Full Name', 'Email Address', 
            'Phone Number', 'Location', 'Budget Range (₹)', 'Signed Status', 'Signature Name', 
            'Word Count', 'Project Details'
        ])

        # Sheet 4: Reviews & Feedback
        ws_reviews = wb.create_sheet(title='Reviews')
        ws_reviews.append(['Review Id', 'Timestamp', 'Author Name', 'Email Address', 'Role', 'Sentiment Keyword', 'Emoji', 'Rating Score', 'Matched Keywords', 'Review Text'])

        for s in reversed(subs):
            sub_id = s.get("id")
            timestamp = s.get("timestamp")
            raw_role = s.get("role", "")
            role_title = "Designer" if raw_role.lower() == "designer" else "Client"
            name = s.get("fullName")
            email = s.get("email")
            phone = s.get("phone")
            location = s.get("location")
            budget = s.get("budget")
            social = s.get("socialHandles", "")
            signed_status = s.get("signed_status", "signed")
            sig_name = s.get("signatureFullName", "")
            word_count = s.get("wordCount", 0)
            desc = s.get("description")

            if role_title == 'Designer':
                ws_designers.append([sub_id, timestamp, name, email, phone, location, budget, signed_status, sig_name, word_count, social, desc])
            else:
                ws_clients.append([sub_id, timestamp, name, email, phone, location, budget, signed_status, sig_name, word_count, desc])

            ws_master.append([sub_id, timestamp, role_title, name, email, phone, location, budget, signed_status, sig_name, word_count, desc])

        for r in reversed(reviews):
            ws_reviews.append([
                r.get("id"), r.get("timestamp"), r.get("name"), r.get("email"),
                r.get("role"), r.get("ratingKeyword"), r.get("emoji"),
                r.get("ratingScore"), r.get("matchedKeywords"), r.get("reviewText")
            ])

        apply_beautification(ws_designers)
        apply_beautification(ws_clients)
        apply_beautification(ws_master)
        apply_beautification(ws_reviews)

        wb.save(EXCEL_PATH)
        print(f"[OK 1s Auto-Sync] Synced ALL {len(subs)} Cumulative Submissions with Signed Status into Excel & SQLite DB.")

    except PermissionError:
        print(f"[Notice] Excel file is open in Microsoft Excel. Saved to SQLite & CSV. Auto-updating Excel when closed...")

if __name__ == "__main__":
    print("=" * 65)
    print("  YOUR INTERIOR DESK - 1-SECOND REAL-TIME EXCEL AUTO-SYNC DAEMON")
    print("  Vercel Submissions API:", VERCEL_API)
    print("  Vercel Reviews API:    ", VERCEL_REVIEWS_API)
    print("  Target SQLite DB:      ", DB_PATH)
    print("  Target Excel:          ", EXCEL_PATH)
    print("  Interval:               EVERY 1 SECOND")
    print("=" * 65)
    print("[*] Running 1-second real-time sync loop... Press Ctrl+C to stop.\n")

    while True:
        sync_1second()
        time.sleep(1)
